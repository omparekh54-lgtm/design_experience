from __future__ import annotations

import io
import os
import re
import time
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Callable

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine, text

load_dotenv()

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/155 Safari/537.36",
    "Accept": "application/json,text/plain,*/*",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": "https://www.nseindia.com/",
}
SOURCE = "NSE_LEGACY_GSM_LOCAL"
MAX_PDF_BYTES = 700_000
MAX_PDF_PAGES = 3


def _engine():
    url = os.getenv("DATABASE_URL") or os.getenv("PMS_DATABASE_URL")
    if not url:
        raise RuntimeError("DATABASE_URL/PMS_DATABASE_URL is required")
    return create_engine(url, pool_pre_ping=True)


def _session() -> requests.Session:
    s = requests.Session()
    s.headers.update(HEADERS)
    for attempt in range(1,4):
        try:
            s.get("https://www.nseindia.com/", timeout=20).raise_for_status()
            return s
        except Exception:
            if attempt==3:
                raise
            time.sleep(2**(attempt-1))
    return s


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    txt = str(value).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y", "%b %d, %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(txt, fmt).date()
        except ValueError:
            pass
    return None


def _attachment_urls(obj) -> list[str]:
    urls: list[str] = []
    def walk(x):
        if isinstance(x, dict):
            for v in x.values(): walk(v)
        elif isinstance(x, list):
            for v in x: walk(v)
        elif isinstance(x, str) and x.startswith("http") and any(ext in x.lower() for ext in (".zip", ".xlsx", ".xls", ".pdf")):
            urls.append(x)
    walk(obj)
    return list(dict.fromkeys(urls))


def _symbols_from_xlsx(data: bytes, known: set[str]) -> set[str]:
    out: set[str] = set()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if not isinstance(value, str): continue
                token = value.strip().upper()
                if token in known:
                    out.add(token)
                else:
                    for part in re.findall(r"[A-Z0-9&.-]{2,25}", token):
                        if part in known: out.add(part)
    return out


def _symbols_from_pdf(data: bytes, known: set[str]) -> set[str]:
    if len(data) > MAX_PDF_BYTES: return set()
    out: set[str] = set()
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        chunks=[]
        for i in range(min(len(reader.pages), MAX_PDF_PAGES)):
            try: chunks.append(reader.pages[i].extract_text() or "")
            except Exception: continue
        text_all="\n".join(chunks)
    except Exception:
        return out
    tokens=set(re.findall(r"[A-Z0-9&.-]{2,25}", text_all.upper()))
    return tokens.intersection(known)


def _extract_attachment(session: requests.Session, url: str, known: set[str]) -> set[str]:
    last=None
    for attempt in range(1,4):
        try:
            r=session.get(url,timeout=40); r.raise_for_status(); data=r.content
            break
        except Exception as exc:
            last=exc
            if attempt==3: raise
            time.sleep(2**(attempt-1))
    else:
        raise last
    lower=url.lower()
    if lower.endswith(".zip") or data[:2]==b"PK":
        out=set()
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names=zf.namelist()
            for name in names:
                if name.lower().endswith(".xlsx"): out |= _symbols_from_xlsx(zf.read(name),known)
            if not out:
                for name in names:
                    if name.lower().endswith(".pdf"): out |= _symbols_from_pdf(zf.read(name),known)
        return out
    if lower.endswith(".xlsx"): return _symbols_from_xlsx(data,known)
    if lower.endswith(".pdf"): return _symbols_from_pdf(data,known)
    return set()


def _subject(row: dict) -> str:
    for key in ("subject", "sub", "circSubject", "circularSubject", "title"):
        val=row.get(key)
        if val: return str(val)
    return ""


def _circular_date(row: dict) -> date | None:
    for key in ("circDate", "date", "circularDate", "displayDate"):
        d=_parse_date(row.get(key))
        if d: return d
    return None


def _fetch_circulars(session: requests.Session, start: date, end: date) -> list[dict]:
    params={"from_date":start.strftime("%d-%m-%Y"),"to_date":end.strftime("%d-%m-%Y"),"sub":"GSM"}
    last=None
    for attempt in range(1,5):
        try:
            r=session.get("https://www.nseindia.com/api/circulars",params=params,timeout=30)
            if r.status_code in (401,403):
                session.get("https://www.nseindia.com/",timeout=20)
                r=session.get("https://www.nseindia.com/api/circulars",params=params,timeout=30)
            r.raise_for_status(); payload=r.json()
            if isinstance(payload,list): return payload
            if isinstance(payload,dict):
                for key in ("data","circulars","records"):
                    if isinstance(payload.get(key),list): return payload[key]
            return []
        except Exception as exc:
            last=exc
            if attempt<4: time.sleep(min(10,2**(attempt-1)))
    assert last is not None
    raise last


def ingest(progress: Callable[[int, int, str], None] | None = None) -> dict:
    engine=_engine()
    try:
        with engine.connect() as conn:
            instrument_rows=conn.execute(text("""
                select id,nse_symbol from pms.instruments
                where exchange='NSE' and nse_symbol is not null
                  and metadata->>'ingestion_duplicate_of' is null
            """)).mappings().all()
        by_symbol={r["nse_symbol"].upper():r["id"] for r in instrument_rows}
        known=set(by_symbol)

        session=_session(); events=defaultdict(list); chunks=[]
        cursor=date(2018,1,1); end_all=date(2023,12,31)
        while cursor<=end_all:
            chunk_end=min(cursor+timedelta(days=89),end_all); chunks.append((cursor,chunk_end)); cursor=chunk_end+timedelta(days=1)

        processed=circular_count=attachment_count=attachment_failures=chunk_failures=unresolved_circulars=0
        unresolved_samples=[]
        for start,end in chunks:
            try:
                rows=_fetch_circulars(session,start,end)
            except Exception as exc:
                chunk_failures+=1
                print(f"\nFAILED GSM circular window {start}..{end}: {type(exc).__name__}",flush=True)
                rows=[]
            for row in rows:
                subject=_subject(row)
                if "GSM" not in subject.upper() and "GRADED SURVEILLANCE" not in subject.upper(): continue
                cdate=_circular_date(row) or start
                urls=sorted(_attachment_urls(row),key=lambda u:0 if any(x in u.lower() for x in (".zip",".xlsx")) else 1)
                symbols=set(); retrieved=0; failed=0
                for url in urls:
                    try:
                        found=_extract_attachment(session,url,known); symbols |= found; attachment_count+=1; retrieved+=1
                        if found and any(x in url.lower() for x in (".zip",".xlsx")): break
                    except Exception:
                        attachment_failures+=1; failed+=1
                # A circular can legitimately have no extractable symbols, but if
                # it advertised attachments and every attachment failed to download
                # we have no evidence about its membership changes and must stop.
                if urls and retrieved == 0 and failed == len(urls):
                    unresolved_circulars += 1
                    if len(unresolved_samples) < 20:
                        unresolved_samples.append({"date":str(cdate),"subject":subject[:160],"attachments":len(urls)})
                    continue
                if not symbols: continue
                circular_count+=1; s=subject.lower()
                inactive=any(x in s for x in ("moving out","move out","removed","removal","exclusion","exit from gsm"))
                for symbol in symbols: events[symbol].append((cdate,not inactive))
            processed+=1
            if progress: progress(processed,len(chunks),f"{start}..{end}")
            time.sleep(0.15)

        if chunk_failures or unresolved_circulars:
            raise RuntimeError(
                f"legacy GSM archive incomplete: {chunk_failures} unresolved circular-window failure(s), "
                f"{unresolved_circulars} circular(s) with all attachments unavailable; samples={unresolved_samples}"
            )

        intervals=[]
        for symbol,evs in events.items():
            evs=sorted(set(evs)); active_from=None
            for d,active in evs:
                if active and active_from is None: active_from=d
                elif not active and active_from is not None:
                    intervals.append((by_symbol[symbol],active_from,d-timedelta(days=1),symbol)); active_from=None
            if active_from is not None: intervals.append((by_symbol[symbol],active_from,date(2024,1,22),symbol))

        with engine.begin() as conn:
            for instrument_id,start,end,symbol in intervals:
                conn.execute(text("""
                    insert into pms.surveillance_history(instrument_id,surveillance_type,stage,effective_from,effective_to,source)
                    values(:instrument_id,'GSM','ACTIVE',:effective_from,:effective_to,:source)
                    on conflict do nothing
                """),{"instrument_id":instrument_id,"effective_from":start,"effective_to":end,"source":SOURCE})

        return {"circulars_used":circular_count,"attachments_parsed":attachment_count,
                "attachment_failures":attachment_failures,"circular_window_failures":chunk_failures,
                "unresolved_circulars":unresolved_circulars,"symbols":len(events),"intervals_written":len(intervals),
                "historical_universe_instruments":len(instrument_rows)}
    finally:
        engine.dispose()


if __name__=="__main__":
    result=ingest(lambda n,t,s:print(f"\rLegacy GSM {n}/{t} {s:28}",end="",flush=True))
    print(); print(result)
